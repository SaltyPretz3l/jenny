"""Rich-file spreadsheet inspect adapter tests."""

from __future__ import annotations

import importlib
import zipfile
from pathlib import Path
from types import SimpleNamespace

import pytest

from sidecar.ai.error_codes import (
    CMP_TOOL_RICH_FILES_DEPENDENCY_MISSING,
    CMP_TOOL_RICH_FILES_UNSUPPORTED,
)
from sidecar.ai.tools.contracts import ToolExecutionFailure
from sidecar.ai.tools.workspace import WorkspaceGuard


def _spreadsheet_module():
    try:
        return importlib.import_module("sidecar.ai.tools.builtins.rich_files.spreadsheet")
    except ModuleNotFoundError as exc:
        pytest.fail(f"spreadsheet inspect adapter is not implemented: {exc}")


def _spreadsheet_tool():
    return _spreadsheet_module().spreadsheet_inspect_tool


def _write_workbook(path: Path, *, many_sheets: int = 0) -> None:
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Visible"
    sheet.append(["Name", "Amount", "Computed"])
    sheet.append(["Alpha", 7, "=SUM(B2:B2)"])
    hidden = workbook.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    hidden["A1"] = "private hidden value"
    very_hidden = workbook.create_sheet("VeryHidden")
    very_hidden.sheet_state = "veryHidden"
    very_hidden["A1"] = "=NOW()"
    for index in range(many_sheets):
        extra = workbook.create_sheet(f"Extra {index + 1}")
        extra["A1"] = f"extra {index + 1}"
    workbook.save(path)


def _write_hidden_first_workbook(path: Path) -> None:
    import openpyxl

    workbook = openpyxl.Workbook()
    hidden = workbook.active
    hidden.title = "Hidden First"
    hidden.sheet_state = "hidden"
    hidden["A1"] = "hidden secret"
    visible = workbook.create_sheet("Visible Second")
    visible["A1"] = "visible sample"
    workbook.save(path)


def _write_hidden_dimensions_workbook(path: Path) -> None:
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet["A1"] = "hidden row value"
    sheet["B1"] = "=NOW()"
    sheet["A2"] = "visible value"
    sheet["B2"] = "hidden column value"
    sheet["C2"] = "later visible column"
    sheet["A3"] = "later visible row"
    sheet["C3"] = "later visible cell"
    sheet.row_dimensions[1].hidden = True
    sheet.column_dimensions["B"].hidden = True
    workbook.save(path)


def _replace_zip_part(path: Path, part_name: str, content: bytes) -> None:
    replacement = path.with_suffix(".replacement")
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(replacement, "w") as target:
        for info in source.infolist():
            target.writestr(info, content if info.filename == part_name else source.read(info))
    replacement.replace(path)


def test_spreadsheet_inspect_returns_workbook_summary_and_samples_visible_rows(
    tmp_path: Path,
) -> None:
    workspace_root = tmp_path / "workspace"
    workspace_root.mkdir()
    source_path = workspace_root / "sample.xlsx"
    _write_workbook(source_path)

    result = _spreadsheet_tool()(
        {
            "path": "sample.xlsx",
            "max_sheets": 5,
            "max_rows_per_sheet": 3,
            "max_columns_per_sheet": 3,
        },
        WorkspaceGuard(str(workspace_root)),
    )

    assert result.success is True
    assert result.generated_artifacts == ()
    assert result.metadata["result_kind"] == "spreadsheet_inspect"
    assert "previews" not in result.metadata
    assert result.metadata["source"]["path"] == "sample.xlsx"
    assert result.metadata["summary"]["sheet_count"] == 3
    assert result.metadata["summary"]["visible_sheet_count"] == 1
    assert result.metadata["summary"]["hidden_sheet_count"] == 2
    assert result.metadata["summary"]["macro_enabled"] is False
    visible = result.metadata["summary"]["sheets"][0]
    assert visible["name"] == "Visible"
    assert visible["state"] == "visible"
    assert visible["has_formulas"] is True
    assert visible["sampled_rows"][0] == ["Name", "Amount", "Computed"]
    assert visible["sampled_rows"][1] == ["Alpha", 7, "[formula]"]
    hidden = result.metadata["summary"]["sheets"][1]
    assert hidden["name"] == "Hidden"
    assert hidden["has_formulas"] is False
    assert hidden["sampled_rows"] == []
    assert hidden["sample_skipped"] == "hidden_sheet"
    assert hidden["formula_scan_skipped"] == "hidden_sheet"
    very_hidden = result.metadata["summary"]["sheets"][2]
    assert very_hidden["state"] == "veryHidden"
    assert very_hidden["has_formulas"] is False
    assert very_hidden["sampled_rows"] == []
    assert very_hidden["formula_scan_skipped"] == "hidden_sheet"
    assert "private hidden value" not in result.output
    assert "=SUM" not in result.output
    assert "=NOW" not in result.output
    assert str(workspace_root) not in result.output


def test_spreadsheet_inspect_prioritizes_visible_sheets_for_sheet_cap(
    tmp_path: Path,
) -> None:
    workspace_root = tmp_path / "workspace"
    workspace_root.mkdir()
    source_path = workspace_root / "hidden-first.xlsx"
    _write_hidden_first_workbook(source_path)

    result = _spreadsheet_tool()(
        {"path": "hidden-first.xlsx", "max_sheets": 1},
        WorkspaceGuard(str(workspace_root)),
    )

    summary = result.metadata["summary"]
    assert summary["sheet_count"] == 2
    assert summary["visible_sheet_count"] == 1
    assert summary["hidden_sheet_count"] == 1
    assert summary["omitted_sheet_count"] == 1
    assert len(summary["sheets"]) == 1
    visible = summary["sheets"][0]
    assert visible["name"] == "Visible Second"
    assert visible["state"] == "visible"
    assert visible["sampled_rows"] == [["visible sample"]]
    assert "hidden secret" not in result.output


def test_spreadsheet_inspect_caps_sheet_rows_and_columns(tmp_path: Path) -> None:
    workspace_root = tmp_path / "workspace"
    workspace_root.mkdir()
    source_path = workspace_root / "many.xlsx"
    _write_workbook(source_path, many_sheets=4)

    result = _spreadsheet_tool()(
        {
            "path": "many.xlsx",
            "max_sheets": 1,
            "max_rows_per_sheet": 1,
            "max_columns_per_sheet": 2,
        },
        WorkspaceGuard(str(workspace_root)),
    )

    summary = result.metadata["summary"]
    assert summary["sheet_count"] == 7
    assert summary["omitted_sheet_count"] == 6
    assert len(summary["sheets"]) == 1
    visible = summary["sheets"][0]
    assert visible["sampled_rows"] == [["Name", "Amount"]]
    assert visible["sample_truncated"] is True
    assert visible["formula_scan_truncated"] is True


def test_spreadsheet_inspect_omits_hidden_rows_and_columns(tmp_path: Path) -> None:
    workspace_root = tmp_path / "workspace"
    workspace_root.mkdir()
    source_path = workspace_root / "hidden-dimensions.xlsx"
    _write_hidden_dimensions_workbook(source_path)

    result = _spreadsheet_tool()(
        {
            "path": "hidden-dimensions.xlsx",
            "max_rows_per_sheet": 2,
            "max_columns_per_sheet": 2,
        },
        WorkspaceGuard(str(workspace_root)),
    )

    sheet = result.metadata["summary"]["sheets"][0]
    assert sheet["sampled_rows"] == [
        ["visible value", "later visible column"],
        ["later visible row", "later visible cell"],
    ]
    assert sheet["has_formulas"] is False
    assert sheet["formula_cells_sampled"] == 0
    assert "hidden row value" not in result.output
    assert "hidden column value" not in result.output


def test_spreadsheet_inspect_rejects_invalid_caps(tmp_path: Path) -> None:
    workspace_root = tmp_path / "workspace"
    workspace_root.mkdir()
    source_path = workspace_root / "sample.xlsx"
    _write_workbook(source_path)

    with pytest.raises(ToolExecutionFailure) as exc_info:
        _spreadsheet_tool()(
            {"path": "sample.xlsx", "max_rows_per_sheet": 0},
            WorkspaceGuard(str(workspace_root)),
        )

    assert exc_info.value.code == CMP_TOOL_RICH_FILES_UNSUPPORTED


def test_spreadsheet_inspect_marks_macro_enabled_extensions(tmp_path: Path) -> None:
    workspace_root = tmp_path / "workspace"
    workspace_root.mkdir()
    source_path = workspace_root / "macro.xlsm"
    _write_workbook(source_path)

    result = _spreadsheet_tool()(
        {"path": "macro.xlsm"},
        WorkspaceGuard(str(workspace_root)),
    )

    assert result.success is True
    assert result.metadata["summary"]["macro_enabled"] is True


def test_spreadsheet_inspect_detects_vba_project_in_renamed_package(tmp_path: Path) -> None:
    workspace_root = tmp_path / "workspace"
    workspace_root.mkdir()
    source_path = workspace_root / "renamed.xlsx"
    _write_workbook(source_path)
    with zipfile.ZipFile(source_path, "a") as archive:
        archive.writestr("XL/VBAPROJECT.BIN", b"vba")

    result = _spreadsheet_tool()(
        {"path": "renamed.xlsx"},
        WorkspaceGuard(str(workspace_root)),
    )

    assert result.metadata["summary"]["macro_enabled"] is True


def test_spreadsheet_lazy_worksheet_parse_failure_returns_unsupported(tmp_path: Path) -> None:
    workspace_root = tmp_path / "workspace"
    workspace_root.mkdir()
    source_path = workspace_root / "malformed-sheet.xlsx"
    _write_workbook(source_path)
    _replace_zip_part(
        source_path,
        "xl/worksheets/sheet1.xml",
        (
            b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
            b'<dimension ref="A1"/><sheetData><row r="1"><c r="A1"/></row>'
        ),
    )

    result = _spreadsheet_tool()(
        {"path": "malformed-sheet.xlsx"},
        WorkspaceGuard(str(workspace_root)),
    )

    assert result.success is False
    assert result.metadata["status"] == "unsupported"
    assert result.metadata["failure"]["reason"] == "spreadsheet_parse_failed"


def test_spreadsheet_inspect_dependency_missing_degrades(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    workspace_root = tmp_path / "workspace"
    workspace_root.mkdir()
    (workspace_root / "sample.xlsx").write_bytes(b"PK\x03\x04")
    module = _spreadsheet_module()

    def _missing_openpyxl(name: str):
        if name == "openpyxl":
            raise ModuleNotFoundError("missing openpyxl")
        return importlib.import_module(name)

    monkeypatch.setattr(module.importlib, "import_module", _missing_openpyxl)

    result = module.spreadsheet_inspect_tool(
        {"path": "sample.xlsx"},
        WorkspaceGuard(str(workspace_root)),
    )

    assert result.success is False
    assert result.metadata["status"] == "unavailable"
    assert result.metadata["failure"]["error_code"] == CMP_TOOL_RICH_FILES_DEPENDENCY_MISSING


def test_spreadsheet_inspect_requires_xml_bomb_protection(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    workspace_root = tmp_path / "workspace"
    workspace_root.mkdir()
    (workspace_root / "sample.xlsx").write_bytes(b"PK\x03\x04")
    module = _spreadsheet_module()

    def _fake_openpyxl(name: str):
        if name == "openpyxl":
            return SimpleNamespace(DEFUSEDXML=False)
        return importlib.import_module(name)

    monkeypatch.setattr(module.importlib, "import_module", _fake_openpyxl)

    result = module.spreadsheet_inspect_tool(
        {"path": "sample.xlsx"},
        WorkspaceGuard(str(workspace_root)),
    )

    assert result.success is False
    assert result.metadata["status"] == "unavailable"
    assert result.metadata["failure"]["error_code"] == CMP_TOOL_RICH_FILES_DEPENDENCY_MISSING
    assert "defusedxml" in result.metadata["failure"]["reason"]


def test_spreadsheet_inspect_unsupported_extension_is_nonfatal(tmp_path: Path) -> None:
    workspace_root = tmp_path / "workspace"
    workspace_root.mkdir()
    (workspace_root / "legacy.xls").write_bytes(b"not an ooxml workbook")

    result = _spreadsheet_tool()(
        {"path": "legacy.xls"},
        WorkspaceGuard(str(workspace_root)),
    )

    assert result.success is False
    assert result.metadata["status"] == "unsupported"
    assert result.metadata["failure"]["reason"] == "spreadsheet_format_unsupported"


def test_spreadsheet_inspect_corrupt_workbook_returns_unsupported(tmp_path: Path) -> None:
    workspace_root = tmp_path / "workspace"
    workspace_root.mkdir()
    (workspace_root / "broken.xlsx").write_bytes(b"not really a workbook")

    result = _spreadsheet_tool()(
        {"path": "broken.xlsx"},
        WorkspaceGuard(str(workspace_root)),
    )

    assert result.success is False
    assert result.metadata["status"] == "unsupported"
    assert result.metadata["failure"]["reason"] == "spreadsheet_parse_failed"



def test_spreadsheet_parses_same_bytes_as_preflight(tmp_path, monkeypatch):
    from sidecar.ai.tools.builtins.rich_files import ooxml
    path = tmp_path / "sample.xlsx"
    _write_workbook(path)
    replacement = tmp_path / "replacement.xlsx"
    import openpyxl

    book = openpyxl.Workbook()
    book.active["A1"] = "replacement data"
    book.save(replacement)
    original = ooxml.preflight_ooxml_archive
    def replace_after_preflight(archive):
        result = original(archive)
        path.write_bytes(replacement.read_bytes())
        return result
    monkeypatch.setattr(ooxml, "preflight_ooxml_archive", replace_after_preflight)
    module = _spreadsheet_module()
    if hasattr(module, "preflight_ooxml_archive"):
        monkeypatch.setattr(module, "preflight_ooxml_archive", replace_after_preflight)
    result = module.spreadsheet_inspect_tool({"path": "sample.xlsx"}, WorkspaceGuard(str(tmp_path)))
    assert result.success is True
    assert result.metadata["summary"]["sheets"][0]["sampled_rows"][0][0] == "Name"
    assert "replacement data" not in result.output


def test_spreadsheet_hidden_prefix_does_not_materialize_rectangle(tmp_path, monkeypatch):
    import openpyxl

    path = tmp_path / "prefix.xlsx"
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.cell(row=500, column=3000, value="visible")
    for row in range(1, 500):
        sheet.row_dimensions[row].hidden = True
    sheet.column_dimensions.group("A", "DKI", hidden=True)
    book.save(path)
    from openpyxl.worksheet._read_only import ReadOnlyWorksheet
    original = ReadOnlyWorksheet.iter_rows
    rectangles = []
    def bounded_rows(self, **kwargs):
        rows = kwargs.get("max_row", 1) - kwargs.get("min_row", 1) + 1
        cols = kwargs.get("max_col", 1) - kwargs.get("min_col", 1) + 1
        rectangles.append(rows * cols)
        if rows * cols > 20_000:
            raise AssertionError("hidden prefix rectangle exceeds traversal budget")
        return original(self, **kwargs)
    monkeypatch.setattr(ReadOnlyWorksheet, "iter_rows", bounded_rows)
    result = _spreadsheet_tool()({"path": "prefix.xlsx", "max_rows_per_sheet": 1,
                                "max_columns_per_sheet": 1}, WorkspaceGuard(str(tmp_path)))
    assert result.success is True
    assert result.metadata["summary"]["sheets"][0]["sampled_rows"] == [["visible"]]
    assert rectangles and max(rectangles) <= 20_000



def test_spreadsheet_parses_the_validated_bytes_without_reopening(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    import hashlib

    module = importlib.import_module("sidecar.ai.tools.builtins.rich_files.spreadsheet")
    workspace_root = tmp_path / "workspace"
    workspace_root.mkdir()
    path = workspace_root / "sample.xlsx"
    _write_workbook(path)
    validated = hashlib.sha256(path.read_bytes()).hexdigest()
    original_validate = module.validate_rich_file_source

    def validate_then_swap(**kwargs):
        source = original_validate(**kwargs)
        # A swap or growth after validation never reaches the parser: it reads
        # the bytes validation hashed, not the file again.
        path.write_bytes(b"swapped after validation" * 4096)
        return source

    monkeypatch.setattr(module, "validate_rich_file_source", validate_then_swap)
    result = module.spreadsheet_inspect_tool({"path": "sample.xlsx"}, WorkspaceGuard(str(workspace_root)))

    assert result.success is True
    assert result.metadata["source"]["sha256"] == validated
    assert "Alpha" in result.output


def test_spreadsheet_dimension_work_is_bounded_before_sampling(tmp_path, monkeypatch):
    module = _spreadsheet_module()
    path = tmp_path / "large.xlsx"
    _write_workbook(path)
    calls = []
    def bounded_dimensions(archive, *, worksheet, element_tree, max_rows, max_columns):
        calls.append((max_rows, max_columns))
        assert max_rows <= 10_000
        assert max_columns <= 16_384
        return set(), set()
    monkeypatch.setattr(module, "_hidden_dimensions", bounded_dimensions)
    sheet = SimpleNamespace(max_row=10**9, max_column=10**9, sheet_state="visible", title="Large",
                            iter_rows=lambda **kwargs: iter([("first",)]))
    payload = module._inspect_sheet(archive=None, element_tree=None, worksheet=sheet,
                                   max_rows=1, max_columns=1)
    assert calls == [(10_000, 16_384)]
    assert payload["sample_truncated"] is True


def test_spreadsheet_sparse_range_work_is_bounded_before_iteration(monkeypatch):
    module = _spreadsheet_module()
    calls = []
    def iterate(**kwargs):
        calls.append(kwargs)
        return iter([])
    worksheet = SimpleNamespace(iter_rows=iterate)
    sampled, formulas, count = module._sample_visible_cells(
        worksheet=worksheet, visible_rows=list(range(1, 201, 2)),
        visible_columns=list(range(1, 101, 2)),
    )
    assert calls == []
    assert sampled == []
    assert formulas is False and count == 0


def _strip_dimension(path: Path) -> None:
    import re

    with zipfile.ZipFile(path) as archive:
        sheet_xml = archive.read("xl/worksheets/sheet1.xml")
    stripped, count = re.subn(rb"<dimension[^>]*/>", b"", sheet_xml)
    assert count == 1
    _replace_zip_part(path, "xl/worksheets/sheet1.xml", stripped)


def test_spreadsheet_without_dimension_is_sized_from_its_cells(tmp_path: Path) -> None:
    workspace_root = tmp_path / "workspace"
    workspace_root.mkdir()
    source_path = workspace_root / "nodim.xlsx"
    _write_workbook(source_path)
    _strip_dimension(source_path)

    result = _spreadsheet_tool()({"path": "nodim.xlsx"}, WorkspaceGuard(str(workspace_root)))

    assert result.success is True
    visible = result.metadata["summary"]["sheets"][0]
    assert visible["dimension_unknown"] is True
    assert visible["max_row"] == 2
    assert visible["max_column"] == 3
    assert visible["sampled_rows"][0] == ["Name", "Amount", "Computed"]
    assert visible["sampled_rows"][1][0] == "Alpha"
    assert visible["has_formulas"] is True
    assert visible["sample_truncated"] is False


def test_spreadsheet_dimension_scan_is_bounded(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    module = _spreadsheet_module()
    workspace_root = tmp_path / "workspace"
    workspace_root.mkdir()
    source_path = workspace_root / "nodim.xlsx"
    _write_workbook(source_path)
    _strip_dimension(source_path)
    monkeypatch.setattr(module, "MAX_SAMPLE_ROW_INDEX", 1)

    result = module.spreadsheet_inspect_tool(
        {"path": "nodim.xlsx"}, WorkspaceGuard(str(workspace_root))
    )

    visible = result.metadata["summary"]["sheets"][0]
    assert visible["dimension_unknown"] is True
    assert visible["max_row"] == 1
    assert visible["sampled_rows"] == [["Name", "Amount", "Computed"]]
    assert visible["sample_truncated"] is True


def test_cell_column_parses_a1_references() -> None:
    module = _spreadsheet_module()
    assert module._cell_column("A1") == 1
    assert module._cell_column("bc12") == 55
    assert module._cell_column("XFD1048576") == 16_384
    assert module._cell_column("") == 0
    assert module._cell_column("12") == 0
